"""
Word 高级操作 MCP 工具模块

提供 Word 文档的读取和高级创建能力，基于 python-docx。

工具列表:
- word_read           读取 Word 文档内容（段落、表格）
- word_create_doc     从零创建 Word 文档（升级版）
- word_add_paragraph  向已有文档追加段落
- word_add_heading    向已有文档追加标题
- word_add_table      向已有文档追加表格
"""

import os
import json
from mcp.server.fastmcp import FastMCP


def _import_docx():
    try:
        from docx import Document as DocxDoc
        from docx.shared import Pt, Inches, RGBColor, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.style import WD_STYLE_TYPE
        return DocxDoc, Pt, Inches, RGBColor, Cm, WD_ALIGN_PARAGRAPH, WD_STYLE_TYPE
    except ImportError:
        return None


def _resolve_safe_path(path: str, base_dir: str) -> str:
    if os.path.isabs(path):
        if not path.startswith(os.path.abspath(base_dir)):
            raise ValueError(f"路径越权: {path}")
        return path
    return os.path.join(base_dir, path)


def _set_chinese_font(run, font_name: str):
    """设置中文字体"""
    try:
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.insert(0, rFonts)
        rFonts.set(qn('w:eastAsia'), font_name)
    except Exception:
        pass


def register_word_tools(mcp: FastMCP, base_dir: str = "/data", **kwargs):
    """注册 Word 高级操作工具"""

    deps = _import_docx()
    if deps is None:
        print("[mcp-file-doc] WARNING: python-docx 未安装，Word 高级工具不可用", flush=True)
        return
    DocxDoc, Pt, Inches, RGBColor, Cm, WD_ALIGN_PARAGRAPH, WD_STYLE_TYPE = deps

    server_base_url = kwargs.get("server_base_url", "")

    ALIGN_MAP = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }

    COLOR_MAP = {
        '黑色': '000000', '黑': '000000', '红色': 'FF0000', '红': 'FF0000',
        '绿色': '00FF00', '绿': '00FF00', '蓝色': '0000FF', '蓝': '0000FF',
        '白色': 'FFFFFF', '白': 'FFFFFF', '灰色': '808080', '灰': '808080',
    }

    def _make_download_url(abs_file_path: str) -> str:
        """如果文件在 base_dir 下，生成可下载 URL"""
        url_base = server_base_url
        if not url_base:
            try:
                ctx = mcp.get_context()
                req = ctx.request_context.request
                if req is not None:
                    host = req.headers.get("x-forwarded-host") or req.headers.get("host", "localhost:8002")
                    proto = req.headers.get("x-forwarded-proto", "http")
                    url_base = f"{proto}://{host}"
            except Exception:
                url_base = "http://localhost:8002"
        if not url_base:
            return ""
        try:
            abs_base = os.path.abspath(base_dir)
            abs_file = os.path.abspath(abs_file_path)
            if abs_file.startswith(abs_base):
                rel = os.path.relpath(abs_file, abs_base).replace("\\", "/")
                return f"{url_base}/download/{rel}"
        except Exception:
            pass
        return ""

    def _resolve_existing_docx(path: str) -> str:
        """解析已存在的 docx 文件路径，支持相对路径自动尝试 uploads/ 子目录"""
        target = _resolve_safe_path(path, base_dir)
        if os.path.isfile(target):
            return target
        # 相对路径：尝试 base_dir/uploads/ 子目录
        if not os.path.isabs(path):
            alt = os.path.join(base_dir, "uploads", os.path.basename(path))
            if os.path.isfile(alt):
                return alt
        return target

    # ==========================================
    # 读取
    # ==========================================

    @mcp.tool()
    async def word_read(path: str) -> str:
        """读取 Word 文档内容，返回段落和表格结构。

        Args:
            path: .docx 文件路径

        Returns:
            JSON: {"status":"ok", "paragraphs":[...], "tables":[...], "paragraph_count":N, "table_count":N}
        """
        try:
            target = _resolve_existing_docx(path)
            if not os.path.isfile(target):
                return json.dumps({"error": f"文件不存在: {path}"}, ensure_ascii=False)

            doc = DocxDoc(target)

            paragraphs = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    style = para.style.name if para.style else "Normal"
                    paragraphs.append({"text": text, "style": style})

            tables = []
            for ti, table in enumerate(doc.tables):
                rows_data = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    rows_data.append(cells)
                tables.append({"index": ti, "rows": len(rows_data), "cols": len(rows_data[0]) if rows_data else 0, "data": rows_data})

            return json.dumps({
                "status": "ok",
                "file": path,
                "paragraph_count": len(paragraphs),
                "table_count": len(tables),
                "paragraphs": paragraphs,
                "tables": tables,
            }, ensure_ascii=False, indent=2)
        except Exception as e:
            return json.dumps({"error": str(e)}, ensure_ascii=False)

    # ==========================================
    # 创建/追加
    # ==========================================

    @mcp.tool()
    async def word_create_doc(
        text: str,
        filename: str = "document.docx",
        font_name: str = "宋体",
        font_size_pt: float = 15.0,
        font_color: str = "000000",
        alignment: str = "center",
        save_to: str = "",
    ) -> str:
        """从零创建一个包含指定文字的 Word 文档（.docx）。

        Args:
            text: 要写入的文本内容
            filename: 输出文件名，默认 "document.docx"
            font_name: 字体，默认 "宋体"
            font_size_pt: 字号（磅），默认 15.0
            font_color: 字体颜色，默认 "000000"（黑色）
            alignment: 对齐方式，"left"/"center"/"right"，默认 "center"
            save_to: 保存路径（可选）

        Returns:
            JSON: {"status":"ok", "output_path":"...", "text":"..."}
        """
        try:
            align = ALIGN_MAP.get(alignment.lower(), WD_ALIGN_PARAGRAPH.CENTER)
            color_hex = COLOR_MAP.get(font_color.strip().lower(), font_color.lstrip('#').upper())

            doc = DocxDoc()
            para = doc.add_paragraph()
            para.alignment = align
            run = para.add_run(text)
            run.font.name = font_name
            run.font.size = Pt(font_size_pt)
            _set_chinese_font(run, font_name)

            try:
                if len(color_hex) == 6:
                    r = int(color_hex[0:2], 16)
                    g = int(color_hex[2:4], 16)
                    b = int(color_hex[4:6], 16)
                    run.font.color.rgb = RGBColor(r, g, b)
            except ValueError:
                pass

            safe_name = os.path.basename(filename)
            if not safe_name.lower().endswith('.docx'):
                safe_name += '.docx'

            if save_to and save_to.strip():
                out_path = _resolve_safe_path(save_to, base_dir)
            else:
                out_path = os.path.join(base_dir, "uploads", safe_name)

            os.makedirs(os.path.dirname(out_path), exist_ok=True)
            doc.save(out_path)

            result = {
                "status": "ok",
                "output_path": out_path,
                "text": text,
                "font": font_name,
                "font_size_pt": font_size_pt,
                "alignment": alignment,
            }

            download_url = _make_download_url(out_path)
            if download_url:
                result["download_url"] = download_url

            return json.dumps(result, ensure_ascii=False, indent=2)
        except Exception as e:
            return json.dumps({"error": str(e)}, ensure_ascii=False)

    @mcp.tool()
    async def word_add_paragraph(
        path: str,
        text: str,
        font_name: str = "宋体",
        font_size_pt: float = 12.0,
        font_color: str = "000000",
        bold: bool = False,
        alignment: str = "left",
    ) -> str:
        """向已有 Word 文档追加段落。

        Args:
            path: .docx 文件路径（必须已存在）
            text: 段落文本
            font_name: 字体，默认 "宋体"
            font_size_pt: 字号（磅），默认 12.0
            font_color: 颜色，默认黑色
            bold: 是否加粗
            alignment: 对齐方式

        Returns:
            JSON: {"status":"ok", "added_text":"..."}
        """
        try:
            target = _resolve_existing_docx(path)
            if not os.path.isfile(target):
                return json.dumps({"error": f"文件不存在: {path}"}, ensure_ascii=False)

            doc = DocxDoc(target)
            align = ALIGN_MAP.get(alignment.lower(), WD_ALIGN_PARAGRAPH.LEFT)

            para = doc.add_paragraph()
            para.alignment = align
            run = para.add_run(text)
            run.font.name = font_name
            run.font.size = Pt(font_size_pt)
            run.bold = bold
            _set_chinese_font(run, font_name)

            color_hex = COLOR_MAP.get(font_color.strip().lower(), font_color.lstrip('#').upper())
            try:
                if len(color_hex) == 6:
                    run.font.color.rgb = RGBColor(
                        int(color_hex[0:2], 16), int(color_hex[2:4], 16), int(color_hex[4:6], 16)
                    )
            except ValueError:
                pass

            doc.save(target)

            return json.dumps({
                "status": "ok",
                "file": path,
                "added_text": text,
                "font": font_name,
                "font_size_pt": font_size_pt,
            }, ensure_ascii=False, indent=2)
        except Exception as e:
            return json.dumps({"error": str(e)}, ensure_ascii=False)

    @mcp.tool()
    async def word_add_heading(
        path: str,
        text: str,
        level: int = 1,
        font_name: str = "黑体",
    ) -> str:
        """向已有 Word 文档追加标题。

        Args:
            path: .docx 文件路径
            text: 标题文本
            level: 标题级别 1-9，默认 1（一级标题）
            font_name: 字体，默认 "黑体"

        Returns:
            JSON: {"status":"ok", "heading_level":N, "added_text":"..."}
        """
        try:
            target = _resolve_existing_docx(path)
            if not os.path.isfile(target):
                return json.dumps({"error": f"文件不存在: {path}"}, ensure_ascii=False)

            doc = DocxDoc(target)
            heading = doc.add_heading(text, level=level)
            _set_chinese_font(heading.runs[0], font_name) if heading.runs else None
            doc.save(target)

            return json.dumps({
                "status": "ok",
                "file": path,
                "heading_level": level,
                "added_text": text,
            }, ensure_ascii=False, indent=2)
        except Exception as e:
            return json.dumps({"error": str(e)}, ensure_ascii=False)

    @mcp.tool()
    async def word_add_table(
        path: str,
        data: str,
        headers: str = "",
        font_name: str = "宋体",
        font_size_pt: float = 11.0,
    ) -> str:
        """向已有 Word 文档追加表格。

        Args:
            path: .docx 文件路径
            data: JSON 字符串，二维数组 '[[\"c1\",\"c2\"],[\"v1\",\"v2\"]]'
            headers: JSON 数组 '[\"列1\",\"列2\"]'（可选表头）
            font_name: 字体，默认 "宋体"
            font_size_pt: 字号（磅），默认 11.0

        Returns:
            JSON: {"status":"ok", "rows":N, "cols":N}
        """
        try:
            target = _resolve_existing_docx(path)
            if not os.path.isfile(target):
                return json.dumps({"error": f"文件不存在: {path}"}, ensure_ascii=False)

            rows = json.loads(data)
            if not isinstance(rows, list) or len(rows) == 0:
                return json.dumps({"error": "data 必须是非空 JSON 数组"}, ensure_ascii=False)

            header_list = json.loads(headers) if headers else []

            doc = DocxDoc(target)

            # 如果有表头，插入表头行
            if header_list:
                all_rows = [header_list] + rows
            else:
                all_rows = rows

            cols = max(len(r) for r in all_rows)
            table = doc.add_table(rows=len(all_rows), cols=cols, style="Table Grid")

            for ri, row_data in enumerate(all_rows):
                for ci, val in enumerate(row_data):
                    cell = table.cell(ri, ci)
                    cell.text = str(val) if val is not None else ""
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.name = font_name
                            run.font.size = Pt(font_size_pt)
                            _set_chinese_font(run, font_name)

            doc.save(target)

            return json.dumps({
                "status": "ok",
                "file": path,
                "rows": len(all_rows),
                "cols": cols,
                "has_headers": bool(header_list),
            }, ensure_ascii=False, indent=2)
        except Exception as e:
            return json.dumps({"error": str(e)}, ensure_ascii=False)

    print("[mcp-file-doc] Word 高级工具已注册", flush=True)
